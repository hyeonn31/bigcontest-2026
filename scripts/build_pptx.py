"""PT_outline.md의 구조를 실제 PowerPoint 파일(.pptx)로 만든다.

python-pptx로 슬라이드를 코드로 생성 — 그래프(pt_*.png)를 그대로 삽입하고,
텍스트는 PT_outline.md의 문장을 그대로 가져다 씀. 완성 후 PowerPoint에서 열어
자유롭게 수정 가능(폰트·색·문구 전부 편집 가능한 진짜 pptx 파일).
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE = r"C:\SH_Antigravity\bigcontest"
FIG = os.path.join(BASE, "figures")
OUT = os.path.join(BASE, "PT_slides.pptx")

BLUE = RGBColor(0x3B, 0x6F, 0xA0)
RED = RGBColor(0xD4, 0x5B, 0x5B)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def textbox(slide, left, top, width, height, text, size=18, bold=False, color=DARK,
            align=PP_ALIGN.LEFT, font=FONT, line_spacing=1.15):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = line_spacing
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
    return tb


def accent_bar(slide, color=BLUE, top=0.0, height=0.12):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(top), prs.slide_width, Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp


def page_number(slide, n):
    textbox(slide, 12.3, 7.05, 0.8, 0.4, str(n), size=11, color=GRAY, align=PP_ALIGN.RIGHT)


def section_title(slide, title, subtitle=None):
    accent_bar(slide, BLUE, top=0, height=0.12)
    textbox(slide, 0.7, 0.4, 11.5, 1.0, title, size=28, bold=True, color=DARK)
    if subtitle:
        textbox(slide, 0.7, 1.15, 11.5, 0.6, subtitle, size=15, color=GRAY)


n = 0

# ---- 표지 ----
s = add_slide(); set_bg(s)
accent_bar(s, BLUE, top=0, height=0.15)
accent_bar(s, RED, top=7.35, height=0.15)
textbox(s, 1, 2.6, 11.3, 1.5, "기후위기 취약상권 영향 분석", size=40, bold=True, color=DARK, align=PP_ALIGN.CENTER)
textbox(s, 1, 3.7, 11.3, 0.8, "서울 강남구 vs 강원 춘천시 — 2025.07~12", size=18, color=GRAY, align=PP_ALIGN.CENTER)
textbox(s, 1, 6.3, 11.3, 0.5, "2026 빅콘테스트 · AI데이터분석 분야", size=14, color=GRAY, align=PP_ALIGN.CENTER)
n += 1; page_number(s, n)

# ---- 1. 문제 제기 ----
s = add_slide(); set_bg(s)
section_title(s, "기후위기가 상권에 영향을 준다는 건 다들 짐작하지만,")
textbox(s, 0.7, 2.0, 11.5, 1.2, "\u201c어떤 상권이 더 취약한가\u201d는 불명확하다", size=26, bold=True, color=BLUE)
textbox(s, 0.7, 3.3, 11.5, 2.5,
        "서울 강남구와 강원 춘천시, 2025년 7월~12월 184일 동안의\n"
        "카드소비 · 유동인구 · 기상 데이터를 결합해 확인했다.",
        size=18, color=DARK, line_spacing=1.4)
n += 1; page_number(s, n)

# ---- 2. 핵심 결론 ----
s = add_slide(); set_bg(s, RGBColor(0xF7, 0xF3, 0xE9))
textbox(s, 1, 1.3, 11.3, 0.6, "핵심 결론", size=16, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
textbox(s, 1, 2.2, 11.3, 3.2,
        "기후위기는 상권을 균일하게 위협하지 않는다.\n\n"
        "외지 방문객 매출 비중이 높은 관광의존형 상권은\n"
        "호우 시 방문객 소비가 지역주민 대비 최대 3배 더 크게 줄어\n"
        "구조적으로 취약하다.",
        size=24, bold=True, color=DARK, align=PP_ALIGN.CENTER, line_spacing=1.4)
n += 1; page_number(s, n)

# ---- 3. 데이터 소개 ----
s = add_slide(); set_bg(s)
section_title(s, "어떤 데이터로 확인했나")
items = [
    ("신한카드 소비 데이터", "184만+180만 행 — 날짜·지역·업종·고객거주지별 결제금액"),
    ("SK텔레콤 유동인구", "좌표 격자 단위 유동인구 (성별·연령·시간대·요일)"),
    ("기상청 ASOS", "서울·춘천 시간자료 → 일자료로 집계 (기온·강수량 등)"),
]
top = 2.2
for name, desc in items:
    accent_bar(s, BLUE, top=top + 0.05, height=0.35)
    tb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(top), Inches(0.08), Inches(1.0))
    tb.fill.solid(); tb.fill.fore_color.rgb = BLUE; tb.line.fill.background()
    textbox(s, 1.0, top, 4.0, 0.5, name, size=18, bold=True, color=DARK)
    textbox(s, 1.0, top + 0.5, 10.5, 0.5, desc, size=14, color=GRAY)
    top += 1.3
textbox(s, 0.7, 6.3, 11.5, 0.6, "→ 날짜×지역 단위로 병합 (368행 = 184일 × 2지역)", size=14, color=BLUE, bold=True)
n += 1; page_number(s, n)

# ---- 4. 분석 방법 ----
s = add_slide(); set_bg(s)
section_title(s, "분석 방법 (간단히)")
textbox(s, 0.7, 2.0, 11.5, 3.5,
        "회귀분석으로 월 · 요일 · 공휴일 효과를 통제하고,\n"
        "비와 폭염의 순수한 효과만 따로 추정했다.\n\n"
        "· 시계열 자기상관 보정 (HAC 표준오차)\n"
        "· 다중검정 보정 (FDR)\n"
        "· 결론이 특정 기준값에 좌우되지 않는지 강건성 점검",
        size=18, color=DARK, line_spacing=1.5)
textbox(s, 0.7, 6.6, 11.5, 0.5, "상세 방법론은 별첨 기타자료(README.md) 참고", size=12, color=GRAY)
n += 1; page_number(s, n)


def image_slide(title, subtitle, img_path, note=None):
    global n
    s = add_slide(); set_bg(s)
    section_title(s, title, subtitle)
    pic_path = os.path.join(FIG, img_path)
    pic = s.shapes.add_picture(pic_path, Inches(2.0), Inches(1.9), height=Inches(5.0))
    if note:
        textbox(s, 0.7, 7.0, 11.5, 0.4, note, size=11, color=GRAY, align=PP_ALIGN.CENTER)
    n += 1; page_number(s, n)
    return s


# ---- 5-1 ----
image_slide("근거 ①", "\u2018지역 전체\u2019로 보면 애매하다", "pt_01_region_overview.png",
            "강남은 신뢰구간이 0을 걸쳐 유의하지 않음 → 고객을 나눠봐야 함")

# ---- 5-2 메인 ----
image_slide("근거 ② — 핵심 발견", "\u2018고객 구성\u2019으로 쪼개면 뚜렷하다", "pt_02_visitor_local_main.png",
            "춘천: 타 시도 방문객 \u221219.1% vs 동일 시도 거주자 \u22126.4% (둘 다 p<0.001, 신뢰구간 안 겹침)")

# ---- 5-3 ----
image_slide("근거 ③", "우연이 아니다 (강건성 점검)", "pt_03_robustness.png",
            "호우 기준 15~40mm 전 구간에서 격차 유지, 전일 지연효과 없음")

# ---- 6. 시사점 ----
s = add_slide(); set_bg(s)
section_title(s, "시사점 / 제언")
sugg = [
    ("지자체 · 관광공사", "외지 방문객 매출 비중만으로 저비용 \u201c상권 기후 취약도\u201d 지표화 가능"),
    ("소상공인", "호우 예보 시 방문객 이탈 방어책 (취소 유연화, 실내 대체 프로그램) — 합리적 추정"),
    ("카드사", "업종·지역별 비 민감도를 소상공인 리스크 상품의 초기 근거로 활용"),
    ("일반화", "\u201c외지 방문객 비중\u201d 지표는 다른 관광도시에도 적용 가능한 틀"),
]
top = 1.9
for title_, desc in sugg:
    textbox(s, 0.7, top, 3.0, 0.6, title_, size=16, bold=True, color=RED)
    textbox(s, 3.8, top, 8.5, 0.9, desc, size=14, color=DARK, line_spacing=1.3)
    top += 1.25
n += 1; page_number(s, n)

# ---- 7. 한계 ----
s = add_slide(); set_bg(s)
section_title(s, "한계 (정직하게 명시)")
textbox(s, 0.7, 2.0, 11.5, 4.0,
        "· 기상 관측소 1곳/도시 — 강남구는 종로구 관측소 대리값\n"
        "· SK 유동인구는 월×요일 평균 패턴, 일별 실측 아님\n"
        "· 표본 지역 2곳뿐 — 단, \u201c외지 방문객 비중\u201d 지표 자체는\n"
        "  다른 지역에도 적용 가능한 일반화된 틀로 제시",
        size=18, color=DARK, line_spacing=1.6)
n += 1; page_number(s, n)

# ---- 8. 백업(부록 표지) ----
s = add_slide(); set_bg(s, RGBColor(0xEF, 0xEF, 0xEF))
textbox(s, 1, 3.2, 11.3, 1.0, "부록 (Q&A 대비 백업 슬라이드)", size=24, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
n += 1; page_number(s, n)

prs.save(OUT)
print("saved:", OUT, "| slides:", n)
