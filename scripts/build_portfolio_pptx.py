"""[11단계 · 발표자료] 포트폴리오(면접·지원서 첨부)용 슬라이드를 만든다.

대회용 build_pptx.py(결론 중심 11장)와 독자가 다르다: 채용 담당자는 결론보다
"문제를 어떻게 풀었고 무엇을 할 줄 아는가"를 보므로 분석 과정·역할·기술을 앞에 둔다.
같은 그래프(figures/pt_*.png)를 재사용하고, 숫자는 REPORT.md·eda_output과 일치시킨다.

입력: figures/pt_01~03_*.png
출력: Portfolio_slides.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from common import BASE

FIG = os.path.join(BASE, "figures")
OUT = os.path.join(BASE, "Portfolio_slides.pptx")
REPO = "github.com/hyeonn31/bigcontest-2026"

BLUE = RGBColor(0x3B, 0x6F, 0xA0)
RED = RGBColor(0xD4, 0x5B, 0x5B)
GREEN = RGBColor(0x2F, 0x7D, 0x5B)
ORANGE = RGBColor(0xA8, 0x55, 0x2F)
DARK = RGBColor(0x1D, 0x1F, 0x23)
GRAY = RGBColor(0x64, 0x6A, 0x73)
LINE = RGBColor(0xE6, 0xE3, 0xDC)
CREAM = RGBColor(0xF4, 0xEF, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def text(s, left, top, width, height, body, size=16, bold=False, color=DARK,
         align=PP_ALIGN.LEFT, spacing=1.2, anchor=MSO_ANCHOR.TOP):
    """여러 줄 텍스트 상자. 줄 앞에 '**'를 붙이면 그 줄만 굵게"""
    tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        line_bold = bold or line.startswith("**")
        p.text = line.removeprefix("**")
        p.alignment = align
        p.line_spacing = spacing
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), line_bold, color, FONT


def box(s, left, top, width, height, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = s.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.12 / min(width, height)  # 모서리 반경을 상자 크기와 무관하게 약 0.12인치로
    shp.shadow.inherit = False
    return shp


def header(s, eyebrow, title):
    text(s, 0.7, 0.45, 11, 0.4, eyebrow, size=12, bold=True, color=GRAY)
    text(s, 0.7, 0.8, 12, 0.8, title, size=28, bold=True)
    text(s, 12.0, 7.0, 0.9, 0.3, str(len(prs.slides)), size=11, color=GRAY, align=PP_ALIGN.RIGHT)


def code_slide(eyebrow, title, panels):
    """코드 슬라이드: panels = [(제목, 한 줄 설명, 코드)] 2개를 좌우로"""
    s = new_slide()
    header(s, eyebrow, title)
    for i, (label, desc, src) in enumerate(panels):
        left = 0.7 + i * 6.05
        text(s, left, 1.75, 5.85, 0.4, label, size=15, bold=True, color=BLUE)
        text(s, left, 2.15, 5.85, 0.6, desc, size=12, color=GRAY, spacing=1.2)
        box(s, left, 2.85, 5.85, 3.4, fill=RGBColor(0x1F, 0x23, 0x2A), line=None)
        tf = s.shapes.add_textbox(Inches(left + 0.2), Inches(3.0), Inches(5.5), Inches(3.2)).text_frame
        tf.word_wrap = False
        for j, line in enumerate(src.splitlines()):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = line if line else " "
            p.line_spacing = 1.1
            for r in p.runs:
                is_comment = line.lstrip().startswith("#")
                r.font.size, r.font.name = Pt(11.5), "Consolas"
                r.font.color.rgb = RGBColor(0x8F, 0xBC, 0x8F) if is_comment else RGBColor(0xE6, 0xE6, 0xE6)
    return s


# ---------------- 표지 ----------------
s = new_slide(CREAM)
text(s, 0.9, 1.0, 11, 0.4, "데이터 분석 포트폴리오  ·  개인 프로젝트", size=14, bold=True, color=GRAY)
text(s, 0.9, 1.9, 11.5, 2.0, "비가 오면, 어떤 상권이\n더 크게 흔들리는가", size=44, bold=True, spacing=1.1)
text(s, 0.9, 4.0, 11.5, 0.9, "카드소비 · 유동인구 · 기상 데이터로 본 기후위기 취약상권 분석\n서울 강남구 vs 강원 춘천시, 2025년 하반기 184일",
     size=17, color=GRAY, spacing=1.4)
text(s, 0.9, 6.3, 11.5, 0.4, f"2026 빅콘테스트 AI데이터분석 분야 출품   |   {REPO}", size=12, color=GRAY)

# ---------------- 프로젝트 개요 ----------------
s = new_slide()
header(s, "OVERVIEW", "프로젝트 개요")
rows = [
    ("질문", "호우·폭염이 상권 소비를 얼마나 줄이는가? 그 충격은 상권마다 어떻게 다른가?"),
    ("기간 · 형태", "2026.09 ~ 10  ·  개인 프로젝트 (2026 빅콘테스트 AI데이터분석 분야)"),
    ("내 역할", "전 과정 단독 수행: 데이터 정제·병합 → 가설·통계 검정 → 강건성 검증 → 시각화·보고서"),
    ("데이터", "신한카드 소비 104만 + 180만 행, SK텔레콤 유동인구, 기상청 ASOS 시간자료"),
    ("방법", "로그선형 회귀(OLS) + HAC 표준오차, FDR 다중검정 보정, 상호작용항, 강건성 점검"),
    ("도구", "Python (pandas · statsmodels · scipy · matplotlib · python-pptx), Git"),
]
top = 1.9
for k, v in rows:
    text(s, 0.7, top, 2.2, 0.6, k, size=15, bold=True, color=BLUE)
    text(s, 2.9, top, 9.8, 0.7, v, size=15, spacing=1.3)
    s.shapes.add_connector(1, Inches(0.7), Inches(top + 0.72), Inches(12.6), Inches(top + 0.72)).line.color.rgb = LINE
    top += 0.86

# ---------------- 핵심 결론 ----------------
s = new_slide()
header(s, "KEY FINDING", "한 문장 결론")
box(s, 0.7, 1.8, 11.9, 1.3, fill=CREAM, line=None)
text(s, 1.0, 1.8, 11.3, 1.3,
     "기후위기는 상권을 균일하게 위협하지 않는다. 외지 방문객 매출 비중이 높은 관광의존형 상권은\n"
     "호우 시 방문객 소비가 지역주민보다 훨씬 크게 줄어 구조적으로 취약하다.",
     size=17, bold=True, spacing=1.35, anchor=MSO_ANCHOR.MIDDLE)
stats = [("−19.1%", "춘천 호우일\n타 시도 방문객 소비", RED),
         ("−6.4%", "같은 날\n동일 시도 거주자 소비", BLUE),
         ("약 3배", "방문객이 더 크게 감소\n(두 값 모두 p<0.001)", DARK)]
for i, (big, small, color) in enumerate(stats):
    left = 0.7 + i * 4.0
    box(s, left, 3.5, 3.75, 2.4)
    text(s, left + 0.35, 3.75, 3.2, 1.0, big, size=44, bold=True, color=color)
    text(s, left + 0.35, 4.85, 3.2, 0.9, small, size=14, color=GRAY, spacing=1.3)
text(s, 0.7, 6.3, 11.9, 0.6,
     "효과 = 같은 달·같은 요일·공휴일 여부를 통제한 뒤 '비 안 온 평범한 날' 대비 소비 변화율.  호우 = 일 강수량 30mm 이상.",
     size=12, color=GRAY)

# ---------------- 데이터 & 전처리 ----------------
s = new_slide()
header(s, "DATA", "단위가 다른 3종 데이터를 '하루 × 지역' 표 하나로")
data = [("신한카드 소비 ①", "104만 행", "일×지역×업종×시간대"),
        ("신한카드 소비 ②", "180만 행", "일×지역×업종×고객 거주지"),
        ("SK텔레콤 유동인구", "월별 3종×6개월", "월×50m 격자"),
        ("기상청 ASOS", "시간자료 2지점", "시간×관측소")]
top = 1.85
for name, size_, unit in data:
    box(s, 0.7, top, 5.6, 0.85)
    text(s, 0.95, top + 0.08, 3.0, 0.4, name, size=14, bold=True)
    text(s, 0.95, top + 0.43, 5.2, 0.4, f"{size_}  ·  {unit}", size=12, color=GRAY)
    top += 1.0
box(s, 0.7, top + 0.1, 5.6, 0.8, fill=CREAM, line=None)
text(s, 0.7, top + 0.1, 5.6, 0.8, "→  merged_daily.csv  (368행 = 184일 × 2지역)", size=15, bold=True,
     color=BLUE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

text(s, 6.9, 1.85, 5.8, 0.5, "정제 과정에서 잡아낸 데이터 문제", size=16, bold=True)
issues = [("SK 12월 파일 중복", "모든 행이 두 번씩 들어 있어 12월 유동인구가 2배로 부풀어 있었음 → 중복 제거 후 재집계"),
          ("강수량 빈칸의 정체", "결측 표기 1,496건이 '비 없음'인지 습도·이슬점 조건으로 간접 검증"),
          ("12/31 기상 누락", "기상청 날짜 기준(01~24시) 때문에 빠진 마지막 날을 추가 수집해 보완"),
          ("세금·미분류 결제 50%", "소수 거액 결제가 일 합계를 좌우 → 이를 뺀 '핵심 소비'를 새로 정의")]
top = 2.45
for t, d in issues:
    text(s, 6.9, top, 5.8, 0.4, "· " + t, size=14, bold=True, color=RED)
    text(s, 7.15, top + 0.38, 5.5, 0.75, d, size=12.5, color=GRAY, spacing=1.25)
    top += 1.15

# ---------------- 핵심 코드 ①: 전처리 ----------------
code_slide("CODE ①  ·  PREPROCESSING", "핵심 코드 — 단위가 다른 데이터를 하나로", [
    ("SK 50m 격자 → 지역",
     "지역명이 없어 좌표가 크게 점프하는 지점을 경계로 분리, 12월 통째 중복 행 제거",
     "# scripts/aggregate_sk_region_month.py\n"
     "X_SPLIT = 983000   # 좌표 점프 지점 = 두 지역 경계\n"
     "\n"
     "def region_of(x):\n"
     "    return '서울 강남구' if x < X_SPLIT else '강원 춘천시'\n"
     "\n"
     "df = df.drop_duplicates()   # 12월 파일 중복 제거\n"
     "df['region'] = df['X_COORD'].apply(region_of)\n"
     "g = (df.groupby(['STD_YM', 'region'])[value_cols]\n"
     "        .sum().reset_index())"),
    ("기상 날짜 기준 · 핵심 소비 · 방문객 구분",
     "기상청 하루(01~24시)에 맞추고, 세금·미분류를 뺀 '핵심 소비'와 방문객 여부를 만듦",
     "# scripts/build_merged_daily.py\n"
     "w['date'] = (w['dt'] - pd.Timedelta(hours=1)).dt.normalize()\n"
     "weather['heavy_rain_day'] = weather['precip_sum'] >= 30\n"
     "\n"
     "# scripts/aggregate_shinhan_daily_region.py\n"
     "CORE_EXCLUDE = ['세금공과금', 'ZZ_나머지']   # 금액의 50%\n"
     "core = df2[~df2['MCT_RY_CD'].isin(CORE_EXCLUDE)]\n"
     "\n"
     "# 가맹점 시도 != 고객 거주 시도 → 타 시도 방문객\n"
     "core['visitor'] = (core['MCT_SGG_CD'].str.split(' ').str[0]\n"
     "                   != core['CLN_SGG_CD'])"),
])

# ---------------- 핵심 코드 ②: 모형과 검정 ----------------
code_slide("CODE ②  ·  MODEL & TESTS", "핵심 코드 — 모든 검정이 공유하는 회귀 함수", [
    ("로그선형 회귀 + HAC 표준오차",
     "월·요일 더미로 계절·요일 효과를 통제하고, 비의 자기상관은 HAC(3일)로 보정",
     "# scripts/common.py\n"
     "def fit_hac(g, y, terms):\n"
     "    X = pd.get_dummies(g[['month', 'weekday']].astype(str),\n"
     "                       drop_first=True, dtype=float)\n"
     "    for c in terms:\n"
     "        X[c] = g[c].astype(float)\n"
     "    X = sm.add_constant(X).loc[y.index]\n"
     "    return sm.OLS(y, X).fit(cov_type='HAC',\n"
     "                            cov_kwds={'maxlags': 3})\n"
     "\n"
     "# 로그 계수 → 변화율(%)\n"
     "to_pct = lambda b: (np.exp(b) - 1) * 100"),
    ("FDR 보정 · 상호작용항",
     "업종 145개 동시 검정은 q값으로 보정, 두 지역의 반응 '차이'는 상호작용항으로 직접 검정",
     "# scripts/industry_analysis.py\n"
     "res.loc[idx, 'rain_heavy_q'] = multipletests(\n"
     "    res.loc[idx, 'rain_heavy_p'], method='fdr_bh')[1]\n"
     "\n"
     "# scripts/region_interaction_test.py  (기준 = 강남)\n"
     "m['region_cc'] = (m['region'] == '강원 춘천시').astype(float)\n"
     "for t in ['rain_light', 'rain_heavy', 'heat_day']:\n"
     "    m[f'region_cc_x_{t}'] = m['region_cc'] * m[t]\n"
     "res = fit_hac(m, np.log(m['s2c_amt']), terms)\n"
     "# → region_cc_x_rain_heavy: +5.7%, p = 0.399"),
])

# ---------------- 분석 과정 ----------------
s = new_slide()
header(s, "PROCESS", "가설을 세우고, 틀리면 좁혔다")
steps = [
    ("첫 탐색: '기후 효과 없음'?", "세금·미분류 거액 결제가 신호를 덮고 있었음 → '핵심 소비'로 재검정하자 춘천 호우 효과가 드러남", "전환", GREEN),
    ("검정 방법 업그레이드", "순열검정의 약점(자기상관·공휴일·다중검정) → 회귀 + HAC 표준오차 + FDR 보정으로 재분석", "보완", GREEN),
    ("가설 A: 방문객 많은 업종일수록 취약", "업종별 방문객 비중 vs 비 효과 상관 ρ = −0.04(춘천), 0.16(강남) → 관계 없음", "기각", ORANGE),
    ("가설 B: 강남은 지하철·지하상가가 방어", "두 지역 차이를 상호작용항으로 직접 검정 → p = 0.40, 지역 간 비교 주장은 하지 않음", "기각", ORANGE),
    ("최종: 범위는 좁게, 근거는 단단하게", "춘천 안에서 방문객 vs 지역주민 차이(신뢰구간 겹치지 않음) → '외지 방문객 매출 비중'을 취약도 지표로 제안", "채택", GREEN),
]
top = 1.8
for i, (title, desc, verdict, vcolor) in enumerate(steps, 1):
    circ = box(s, 0.7, top + 0.05, 0.5, 0.5, fill=DARK, line=None, shape=MSO_SHAPE.OVAL)
    circ.text_frame.text = str(i)
    r = circ.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(14), True, WHITE, FONT
    circ.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    tag = box(s, 1.5, top + 0.08, 0.8, 0.42, fill=WHITE, line=vcolor)
    tag.text_frame.text = verdict
    r = tag.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(12), True, vcolor, FONT
    tag.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, 2.5, top, 10.2, 0.45, title, size=16, bold=True)
    text(s, 2.5, top + 0.45, 10.2, 0.5, desc, size=13, color=GRAY)
    top += 1.02

# ---------------- 결과 ①: 지역 전체 ----------------
s = new_slide()
header(s, "RESULT ①", "지역 전체로 보면, 결론이 애매하다")
s.shapes.add_picture(os.path.join(FIG, "pt_01_region_overview.png"), Inches(0.7), Inches(1.75), height=Inches(5.2))
text(s, 7.9, 2.2, 4.8, 4.5,
     "**호우일 핵심 소비 변화\n"
     "· 강남 −11.0%  (p = 0.069)\n"
     "   95% 신뢰구간이 0을 포함 → 불확실\n"
     "· 춘천 −8.6%  (p < 0.001)\n"
     "\n"
     "**두 지역 차이를 직접 검정하면\n"
     "· 상호작용항 p = 0.40 → 차이 근거 없음\n"
     "\n"
     "→ '어느 지역이 더 취약한가'만으로는\n"
     "   답할 수 없다. 고객을 나눠 봐야 한다.",
     size=15, spacing=1.35)

# ---------------- 결과 ②: 핵심 발견 ----------------
s = new_slide()
header(s, "RESULT ②  ·  핵심 발견", "고객 구성으로 나누면, 뚜렷하다")
s.shapes.add_picture(os.path.join(FIG, "pt_02_visitor_local_main.png"), Inches(0.5), Inches(1.75), height=Inches(5.0))
text(s, 8.2, 2.2, 4.6, 4.6,
     "**춘천, 호우일 (18일)\n"
     "· 타 시도 방문객  −19.1%\n"
     "   95% CI  −25.0% ~ −12.7%\n"
     "· 동일 시도 거주자  −6.4%\n"
     "   95% CI  −9.9% ~ −2.7%\n"
     "\n"
     "두 신뢰구간이 겹치지 않음\n"
     "→ 방문객이 확실히 더 크게 줄어든다\n"
     "\n"
     "강남은 두 그룹 모두 유의하지 않음",
     size=15, spacing=1.35)

# ---------------- 강건성 ----------------
s = new_slide()
header(s, "ROBUSTNESS", "'30mm를 골라서 나온 우연' 아닌가?")
s.shapes.add_picture(os.path.join(FIG, "pt_03_robustness.png"), Inches(0.6), Inches(1.75), height=Inches(5.0))
checks = [("호우 기준값 15~40mm", "6단계 모두 방문객(−17.6~−19.5%)이 지역주민(−5.9~−9.0%)보다 크게 감소"),
          ("전일 강수 지연효과", "'어제 호우' 변수 추가 → 지연효과 없음(p=0.46~0.73), 당일 효과 유지"),
          ("폭염 기준 30~35℃", "전 구간 유의하지 않음 → 원 결론과 일치")]
top = 2.0
for t, d in checks:
    text(s, 8.1, top, 4.7, 0.4, "✓  " + t, size=15, bold=True, color=GREEN)
    text(s, 8.45, top + 0.42, 4.35, 0.9, d, size=13, color=GRAY, spacing=1.3)
    top += 1.45

# ---------------- 시사점 & 한계 ----------------
s = new_slide()
header(s, "IMPACT & LIMITS", "어디에 쓸 수 있고, 무엇을 말하지 않는가")
text(s, 0.7, 1.8, 6.0, 0.4, "시사점", size=17, bold=True, color=BLUE)
impact = [("지자체 · 관광공사", "고객 거주지 비중만으로 설문 없이 '상권 기후 취약도' 저비용 모니터링"),
          ("카드사 · 금융", "업종·지역별 비 민감도를 날씨 연동 소상공인 상품의 초기 리스크 근거로"),
          ("일반화", "'외지 방문객 매출 비중' 지표는 다른 관광도시에도 같은 방법으로 적용 가능")]
top = 2.35
for t, d in impact:
    text(s, 0.7, top, 6.0, 0.4, t, size=14, bold=True)
    text(s, 0.7, top + 0.38, 5.9, 0.8, d, size=13, color=GRAY, spacing=1.3)
    top += 1.25
box(s, 7.1, 1.8, 5.6, 4.9, fill=CREAM, line=None)
text(s, 7.4, 2.0, 5.1, 0.4, "한계", size=17, bold=True, color=RED)
text(s, 7.4, 2.55, 5.1, 4.0,
     "· 기상 관측소 도시당 1곳 — 강남구는 약 10km 떨어진 종로구 관측소 값으로 대체\n\n"
     "· SK 유동인구는 월×요일 평균 패턴 → 검정에는 쓰지 않고 배경 지표로만 사용\n\n"
     "· 비교 지역 2곳뿐 → 지역 간 비교의 검정력 한계\n\n"
     "· 연관성 분석이며 인과관계를 확정한 것은 아님",
     size=13.5, spacing=1.3)

# ---------------- 기술 & 배운 점 ----------------
s = new_slide()
header(s, "SKILLS & TAKEAWAYS", "사용한 기술과 얻은 교훈")
skills = [("전처리", "인코딩 변환, 좌표 격자 → 행정구역 공간 집계, 다중 소스 병합"),
          ("EDA", "결측·중복·이상치 점검, 치우친 분포의 로그 변환"),
          ("가설검정", "OLS 로그선형 회귀 + HAC, 순열검정, 스피어만 상관"),
          ("보정 · 비교", "FDR(Benjamini–Hochberg), 상호작용항, 95% 신뢰구간"),
          ("강건성", "임계값 민감도 분석, 시차(lag) 효과 검정"),
          ("재현성", "공통 모듈화, 상대경로, 코드로 그래프·슬라이드 생성, Git")]
top = 1.85
for k, v in skills:
    text(s, 0.7, top, 1.8, 0.4, k, size=13.5, bold=True, color=BLUE)
    text(s, 2.5, top, 4.2, 0.75, v, size=13, spacing=1.25)
    top += 0.8
box(s, 7.1, 1.8, 5.6, 4.9, fill=CREAM, line=None)
text(s, 7.4, 2.0, 5.1, 0.4, "얻은 교훈", size=17, bold=True)
text(s, 7.4, 2.55, 5.1, 4.0,
     "**'효과 없음'도 의심한다\n"
     "첫 결과를 그대로 믿었다면 핵심 발견을 놓쳤다. 데이터 구성을 들여다본 것이 결론을 바꿨다.\n\n"
     "**그럴듯한 이야기보다 직접 검정\n"
     "지역별 결과를 따로 읽고 내린 해석이, 차이 자체를 검정하자 근거가 없었다.\n\n"
     "**주장의 크기를 근거에 맞춘다\n"
     "범위를 좁힌 주장이 더 설득력 있었다.",
     size=13.5, spacing=1.3)
text(s, 0.7, 6.85, 11, 0.35, f"코드 · 상세 보고서: {REPO}", size=12, color=GRAY)

prs.save(OUT)
print("saved:", OUT, "| slides:", len(prs.slides))
